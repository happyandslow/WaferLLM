#!/usr/bin/env python3
"""Generate GQA Implementation Comparison slides with white background."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Colors
BLACK = RGBColor(0x1A, 0x1A, 0x2E)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MED_GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF0, 0xF0, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_BLUE = RGBColor(0x2B, 0x6C, 0xB3)
ACCENT_GREEN = RGBColor(0x27, 0x8B, 0x5B)
ACCENT_RED = RGBColor(0xC0, 0x39, 0x2B)
ACCENT_ORANGE = RGBColor(0xE6, 0x7E, 0x22)
TABLE_HEADER_BG = RGBColor(0x2B, 0x6C, 0xB3)
TABLE_ALT_BG = RGBColor(0xF5, 0xF8, 0xFC)
TABLE_BORDER = RGBColor(0xD0, 0xD0, 0xD0)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height


def add_blank_slide():
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    # white background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    return slide


def add_textbox(slide, left, top, width, height, text, font_size=18,
                bold=False, color=DARK_GRAY, alignment=PP_ALIGN.LEFT,
                font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_rich_textbox(slide, left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf


def add_paragraph(tf, text, font_size=18, bold=False, color=DARK_GRAY,
                  alignment=PP_ALIGN.LEFT, space_after=Pt(6), font_name="Calibri",
                  space_before=Pt(0)):
    if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == "":
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    p.space_after = space_after
    p.space_before = space_before
    return p


def add_bullet(tf, text, font_size=16, color=DARK_GRAY, level=0, bold=False,
               space_after=Pt(4)):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Calibri"
    p.level = level
    p.space_after = space_after
    return p


def add_accent_line(slide, left, top, width, color=ACCENT_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_table(slide, rows, cols, left, top, width, height):
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    return table


def style_table_cell(cell, text, font_size=12, bold=False, color=DARK_GRAY,
                     bg_color=None, alignment=PP_ALIGN.LEFT):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = alignment
    cell.text_frame.word_wrap = True
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color


def add_slide_number(slide, num):
    add_textbox(slide, Inches(12.5), Inches(7.05), Inches(0.7), Inches(0.35),
                str(num), font_size=11, color=MED_GRAY, alignment=PP_ALIGN.RIGHT)


def add_section_badge(slide, text, color=ACCENT_BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.6), Inches(0.4), Inches(2.8), Inches(0.45)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    tf.word_wrap = False


# ─────────────────────────────────────────────
# SLIDE 1: Title
# ─────────────────────────────────────────────
slide = add_blank_slide()

# Accent bar at top
add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_BLUE)

add_textbox(slide, Inches(1), Inches(1.8), Inches(11.3), Inches(1.2),
            "GQA Implementation Comparison",
            font_size=44, bold=True, color=BLACK, alignment=PP_ALIGN.CENTER)

add_accent_line(slide, Inches(4.5), Inches(3.1), Inches(4.3), ACCENT_BLUE)

add_textbox(slide, Inches(1), Inches(3.5), Inches(11.3), Inches(0.8),
            "WSE-3-GQA  vs  WSE-3-GQA-cheng",
            font_size=28, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER,
            font_name="Calibri")

add_textbox(slide, Inches(1), Inches(4.5), Inches(11.3), Inches(0.6),
            "Two approaches to Grouped Query Attention on Cerebras WSE-3",
            font_size=18, color=MED_GRAY, alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(1), Inches(5.8), Inches(11.3), Inches(0.5),
            "WaferLLM  |  Decode Module",
            font_size=16, color=MED_GRAY, alignment=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
# SLIDE 2: Background / Problem Statement
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "BACKGROUND")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Why GQA on WaferLLM?",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

tf = add_rich_textbox(slide, Inches(0.6), Inches(1.9), Inches(5.8), Inches(4.5))
add_paragraph(tf, "Multi-Head Attention (MHA)", font_size=20, bold=True, color=ACCENT_BLUE)
add_bullet(tf, "Q, K, V all have the same number of heads (n_heads)", font_size=16)
add_bullet(tf, "KV cache per PE: [dim_p_pe x seq_len_p_pe] -- scales linearly with n_heads", font_size=16)
add_bullet(tf, "Memory-bound at long sequence lengths", font_size=16)

add_paragraph(tf, "", font_size=10)
add_paragraph(tf, "Grouped Query Attention (GQA)", font_size=20, bold=True, color=ACCENT_GREEN)
add_bullet(tf, "Fewer KV heads than Q heads (e.g., Llama-8B: 32 Q, 8 KV)", font_size=16)
add_bullet(tf, "Multiple Q heads share one KV head (4:1 ratio)", font_size=16)
add_bullet(tf, "Reduces KV cache memory and communication volume", font_size=16)

add_paragraph(tf, "", font_size=10)
add_paragraph(tf, "Challenge on WSE-3", font_size=20, bold=True, color=ACCENT_RED)
add_bullet(tf, "Data sharded across P x P PE mesh (e.g., 256 x 256)", font_size=16)
add_bullet(tf, "Communication patterns tightly coupled to tensor dimensions", font_size=16)
add_bullet(tf, "How to map asymmetric Q/KV heads to symmetric mesh?", font_size=16)

# Right side: diagram placeholder
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(7.0), Inches(1.9), Inches(5.8), Inches(4.5)
)
shape.fill.solid()
shape.fill.fore_color.rgb = LIGHT_GRAY
shape.line.color.rgb = TABLE_BORDER

tf2 = shape.text_frame
tf2.word_wrap = True
add_paragraph(tf2, "MHA (Base)", font_size=18, bold=True, color=BLACK,
              alignment=PP_ALIGN.CENTER, space_after=Pt(4))
add_paragraph(tf2, "Q:[dim]  K:[dim]  V:[dim]", font_size=16, color=DARK_GRAY,
              alignment=PP_ALIGN.CENTER, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf2, "n_heads = n_kv_heads = 32", font_size=14, color=MED_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(16))

add_paragraph(tf2, "GQA (Target)", font_size=18, bold=True, color=ACCENT_GREEN,
              alignment=PP_ALIGN.CENTER, space_after=Pt(4))
add_paragraph(tf2, "Q:[dim]  K:[kv_dim]  V:[kv_dim]", font_size=16, color=DARK_GRAY,
              alignment=PP_ALIGN.CENTER, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf2, "n_heads=32, n_kv_heads=8 (4:1 ratio)", font_size=14, color=MED_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(16))

add_paragraph(tf2, "Two design strategies:", font_size=16, bold=True, color=BLACK,
              alignment=PP_ALIGN.CENTER, space_after=Pt(4))
add_paragraph(tf2, "1. Dimensional: reshape tensors (smaller K/V dims)", font_size=14,
              color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER, space_after=Pt(2))
add_paragraph(tf2, "2. Replicated: keep dims, offset into copies", font_size=14,
              color=ACCENT_ORANGE, alignment=PP_ALIGN.CENTER)

add_slide_number(slide, 2)


# ─────────────────────────────────────────────
# SLIDE 3: Executive Summary Table
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "EXECUTIVE SUMMARY")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "At a Glance: Three-Way Comparison",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

rows_data = [
    ["Metric", "WSE-3 (Base/MHA)", "WSE-3-GQA", "WSE-3-GQA-cheng"],
    ["GQA strategy", "None (MHA)", "Dimensional\n(asymmetric K/V dims)", "Replicated-heads\n(runtime offset)"],
    ["Total CSL delta", "Baseline (2519 lines)", "+1177 lines (+47%)", "+35 lines (+1%)"],
    ["K/V weight shape", "[dim x dim]", "[dim x kv_dim]  (4x smaller)", "[n_kv x dim x dim]  (8x larger)"],
    ["QKV tile layout", "Symmetric [Q|K|V]", "Asymmetric [Q|K|V]", "Symmetric [Q|K|V]  (same)"],
    ["Score tensor", "[bsz x seq_len]", "[groups x bsz x seq_len]", "[bsz x seq_len]  (same)"],
    ["GEMV variants", "1", "3 (+2 new)", "2 (+1 new)"],
    ["Comm reduce fns", "7", "~12 (+5 new)", "~8 (+1 new)"],
    ["RoPE", "Enabled", "Commented out", "Enabled"],
    ["SILU bug", "Present", "Inherited", "Fixed"],
]

table = add_table(slide, len(rows_data), 4,
                  Inches(0.6), Inches(1.85), Inches(12.1), Inches(5.0))

col_widths = [Inches(2.2), Inches(3.0), Inches(3.5), Inches(3.4)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for row_idx, row_data in enumerate(rows_data):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        if row_idx == 0:
            style_table_cell(cell, cell_text, font_size=13, bold=True,
                             color=WHITE, bg_color=TABLE_HEADER_BG,
                             alignment=PP_ALIGN.CENTER)
        else:
            bg = TABLE_ALT_BG if row_idx % 2 == 0 else WHITE
            is_bold = col_idx == 0
            clr = BLACK if col_idx == 0 else DARK_GRAY
            style_table_cell(cell, cell_text, font_size=12, bold=is_bold,
                             color=clr, bg_color=bg)

add_slide_number(slide, 3)


# ─────────────────────────────────────────────
# SLIDE 4: Design Philosophy
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "DESIGN PHILOSOPHY")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Dimensional vs Replicated GQA",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

# Left column - Dimensional
shape_l = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6), Inches(1.9), Inches(5.8), Inches(5.0)
)
shape_l.fill.solid()
shape_l.fill.fore_color.rgb = RGBColor(0xEE, 0xF4, 0xFB)
shape_l.line.color.rgb = ACCENT_BLUE

tf = shape_l.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA: Dimensional", font_size=22, bold=True, color=ACCENT_BLUE,
              alignment=PP_ALIGN.CENTER, space_after=Pt(8))
add_paragraph(tf, "Compile-time dimension reshaping", font_size=14, color=MED_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(12))

add_paragraph(tf, "kv_dim_p_pe = (n_kv_heads * head_dim) / P", font_size=13,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "gqa_group_size = dim_p_pe / kv_dim_p_pe", font_size=13,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(10))

add_paragraph(tf, "Llama-8B (P=256):", font_size=14, bold=True, color=BLACK, space_after=Pt(4))
add_paragraph(tf, "  Q:[16x16]  K:[16x4]  V:[16x4]", font_size=13,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "  QKV tile: [16 + 2*4 = 24]  (was 48)", font_size=13,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "  Score: [4 x bsz x seq_len]", font_size=13,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(12))

add_paragraph(tf, "Pros", font_size=16, bold=True, color=ACCENT_GREEN, space_after=Pt(4))
add_bullet(tf, "K/V memory shrinks by GQA ratio (4x for 4:1)", font_size=13, color=DARK_GRAY)
add_bullet(tf, "Less communication volume (asymmetric QKV reduce)", font_size=13, color=DARK_GRAY)
add_bullet(tf, "Includes KV cache mgmt + variable-length attention", font_size=13, color=DARK_GRAY)

add_paragraph(tf, "Cons", font_size=16, bold=True, color=ACCENT_RED, space_after=Pt(4))
add_bullet(tf, "+47% CSL code -- heavy refactor", font_size=13, color=DARK_GRAY)
add_bullet(tf, "New axis mode, new GEMV variants, new comm fns", font_size=13, color=DARK_GRAY)
add_bullet(tf, "RoPE commented out (strided-DSD bug)", font_size=13, color=DARK_GRAY)

# Right column - Replicated
shape_r = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(6.9), Inches(1.9), Inches(5.8), Inches(5.0)
)
shape_r.fill.solid()
shape_r.fill.fore_color.rgb = RGBColor(0xFD, 0xF2, 0xE9)
shape_r.line.color.rgb = ACCENT_ORANGE

tf = shape_r.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA-cheng: Replicated", font_size=22, bold=True, color=ACCENT_ORANGE,
              alignment=PP_ALIGN.CENTER, space_after=Pt(8))
add_paragraph(tf, "Runtime offset into replicated heads", font_size=14, color=MED_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(12))

add_paragraph(tf, "kv_head_id = quotient_x % n_kv_heads", font_size=13,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "k_offset = kv_head_id * dim_p_pe * dim_p_pe", font_size=13,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(10))

add_paragraph(tf, "Llama-8B (P=256):", font_size=14, bold=True, color=BLACK, space_after=Pt(4))
add_paragraph(tf, "  Q:[16x16]  K:[8*16x16]  V:[8*16x16]", font_size=13,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "  QKV tile: [3*16 = 48]  (same as base)", font_size=13,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "  Score: [bsz x seq_len]  (same as base)", font_size=13,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(12))

add_paragraph(tf, "Pros", font_size=16, bold=True, color=ACCENT_GREEN, space_after=Pt(4))
add_bullet(tf, "Only +1% CSL code -- minimal structural changes", font_size=13, color=DARK_GRAY)
add_bullet(tf, "RoPE preserved, SILU bug fixed", font_size=13, color=DARK_GRAY)
add_bullet(tf, "Comm library essentially unchanged", font_size=13, color=DARK_GRAY)

add_paragraph(tf, "Cons", font_size=16, bold=True, color=ACCENT_RED, space_after=Pt(4))
add_bullet(tf, "K/V memory grows by n_kv_heads (8x for 8 heads)", font_size=13, color=DARK_GRAY)
add_bullet(tf, "Same communication volume as MHA (no reduction)", font_size=13, color=DARK_GRAY)
add_bullet(tf, "No KV cache mgmt or variable-length attention", font_size=13, color=DARK_GRAY)

add_slide_number(slide, 4)


# ─────────────────────────────────────────────
# SLIDE 5: Per-PE Memory Comparison
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "MEMORY ANALYSIS")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Per-PE Memory Comparison (Llama-8B, P=256, 4:1 GQA)",
            font_size=30, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

mem_rows = [
    ["Buffer", "Base (MHA)", "WSE-3-GQA", "GQA-cheng", "GQA vs Base", "cheng vs Base"],
    ["Q weight", "256 f16", "256 f16", "256 f16", "--", "--"],
    ["K weight", "256 f16", "64 f16", "2048 f16", "4x smaller", "8x larger"],
    ["V weight", "256 f16", "64 f16", "2048 f16", "4x smaller", "8x larger"],
    ["KV cache (K)", "320 f16", "80 f16", "2560 f16", "4x smaller", "8x larger"],
    ["KV cache (V)", "320 f16", "80 f16", "2560 f16", "4x smaller", "8x larger"],
    ["Score", "20 f16", "80 f16", "20 f16", "4x larger", "--"],
    ["QKV tile", "48 f16", "24 f16", "48 f16", "2x smaller", "--"],
]

table = add_table(slide, len(mem_rows), 6,
                  Inches(0.6), Inches(1.85), Inches(12.1), Inches(3.5))

col_widths = [Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.8), Inches(2.45), Inches(2.45)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for row_idx, row_data in enumerate(mem_rows):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        if row_idx == 0:
            style_table_cell(cell, cell_text, font_size=13, bold=True,
                             color=WHITE, bg_color=TABLE_HEADER_BG,
                             alignment=PP_ALIGN.CENTER)
        else:
            bg = TABLE_ALT_BG if row_idx % 2 == 0 else WHITE
            is_bold = col_idx == 0
            clr = BLACK if col_idx == 0 else DARK_GRAY
            # Color the comparison columns
            if col_idx == 4 and "smaller" in cell_text:
                clr = ACCENT_GREEN
            elif col_idx == 4 and "larger" in cell_text:
                clr = ACCENT_RED
            elif col_idx == 5 and "larger" in cell_text:
                clr = ACCENT_RED
            style_table_cell(cell, cell_text, font_size=12, bold=is_bold or (col_idx >= 4 and cell_text != "--"),
                             color=clr, bg_color=bg, alignment=PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT)

# Key insight box
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6), Inches(5.6), Inches(12.1), Inches(1.3)
)
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xFD, 0xF9, 0xED)
shape.line.color.rgb = ACCENT_ORANGE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "  Key Insight", font_size=16, bold=True, color=ACCENT_ORANGE, space_after=Pt(4))
add_paragraph(tf, "  WSE-3-GQA reduces K/V memory by the GQA ratio (4x for 4:1) but expands score buffers.",
              font_size=14, color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "  WSE-3-GQA-cheng increases K/V memory by n_kv_heads (8x) because every PE stores all heads.",
              font_size=14, color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "  For large models with many KV heads, the dimensional approach saves significant per-PE memory.",
              font_size=14, color=DARK_GRAY)

add_slide_number(slide, 5)


# ─────────────────────────────────────────────
# SLIDE 6: Code Changes -- QKV Projection
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "CODE CHANGES")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "QKV Projection & Tile Layout",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

# Left: Dimensional
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6), Inches(1.9), Inches(5.8), Inches(4.8)
)
shape.fill.solid()
shape.fill.fore_color.rgb = LIGHT_GRAY
shape.line.color.rgb = ACCENT_BLUE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA: Asymmetric Reshape", font_size=18, bold=True,
              color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER, space_after=Pt(10))

add_paragraph(tf, "// QKV tile restructured:", font_size=12, color=MED_GRAY,
              font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "- var QKV_tile: [bsz * 3 * dim_p_pe]f16;", font_size=12,
              color=ACCENT_RED, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "+ var QKV_tile: [bsz*(dim_p_pe+2*kv_dim_p_pe)]f16;", font_size=12,
              color=ACCENT_GREEN, font_name="Consolas", space_after=Pt(10))

add_paragraph(tf, "// K projection output shrinks:", font_size=12, color=MED_GRAY,
              font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "- Kt = dim_p_pe;  Nt = dim_p_pe;   // 16->16", font_size=12,
              color=ACCENT_RED, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "+ Kt = dim_p_pe;  Nt = kv_dim_p_pe; // 16->4", font_size=12,
              color=ACCENT_GREEN, font_name="Consolas", space_after=Pt(10))

add_paragraph(tf, "// QKV reduce volume:", font_size=12, color=MED_GRAY,
              font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "  3*dim_p_pe  ->  dim_p_pe + 2*kv_dim_p_pe", font_size=12,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "  48 -> 24  (50% less communication)", font_size=13,
              bold=True, color=ACCENT_GREEN, font_name="Consolas")

# Right: Replicated
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(6.9), Inches(1.9), Inches(5.8), Inches(4.8)
)
shape.fill.solid()
shape.fill.fore_color.rgb = LIGHT_GRAY
shape.line.color.rgb = ACCENT_ORANGE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA-cheng: Same Layout, Offset", font_size=18, bold=True,
              color=ACCENT_ORANGE, alignment=PP_ALIGN.CENTER, space_after=Pt(10))

add_paragraph(tf, "// QKV tile UNCHANGED:", font_size=12, color=MED_GRAY,
              font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "  var QKV_tile: [bsz * 3 * dim_p_pe]f16;", font_size=12,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(10))

add_paragraph(tf, "// K projection -- only add offset:", font_size=12, color=MED_GRAY,
              font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "+ var k_off = kv_head_id*dim*dim;", font_size=12,
              color=ACCENT_GREEN, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "+ right_dsd = @increment_dsd_offset(", font_size=12,
              color=ACCENT_GREEN, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "+     right_dsd, k_off, f16);", font_size=12,
              color=ACCENT_GREEN, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "  Kt=dim_p_pe; Nt=dim_p_pe; // same dims", font_size=12,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(10))

add_paragraph(tf, "// QKV reduce volume:", font_size=12, color=MED_GRAY,
              font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "  3*dim_p_pe (unchanged from base)", font_size=12,
              color=DARK_GRAY, font_name="Consolas", space_after=Pt(2))
add_paragraph(tf, "  48 -> 48  (no reduction in comm)", font_size=13,
              bold=True, color=ACCENT_RED, font_name="Consolas")

add_slide_number(slide, 6)


# ─────────────────────────────────────────────
# SLIDE 7: Attention Score Computation
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "CODE CHANGES")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Attention Score Computation",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

# Left: Dimensional
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6), Inches(1.9), Inches(5.8), Inches(5.0)
)
shape.fill.solid()
shape.fill.fore_color.rgb = LIGHT_GRAY
shape.line.color.rgb = ACCENT_BLUE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA (+80 lines)", font_size=18, bold=True,
              color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER, space_after=Pt(8))

add_paragraph(tf, "Multi-group, variable-length, head-scoped:", font_size=14,
              bold=True, color=DARK_GRAY, space_after=Pt(6))

add_paragraph(tf, "fn score_matvec_mult() {", font_size=12,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "  for (b in bsz) {", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "    for (g in gqa_group_size) {", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "      score[g,b,:] = Q[b,g,:] @ K[b]^T", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "    }", font_size=12, font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "  }", font_size=12, font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "  Kt=kv_dim; Nt=iter_num; // var-len", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "  vecmat_var_right_2d_strided();", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "  all_reduce_kv_head_scoped(); //NEW", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "}", font_size=12, font_name="Consolas", color=DARK_GRAY, space_after=Pt(8))

add_paragraph(tf, "New: nested loops, 2 GEMV variants,", font_size=13,
              bold=True, color=ACCENT_RED, space_after=Pt(2))
add_paragraph(tf, "new axis=3 comm mode, head-scoped reduce", font_size=13,
              bold=True, color=ACCENT_RED)

# Right: Replicated
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(6.9), Inches(1.9), Inches(5.8), Inches(5.0)
)
shape.fill.solid()
shape.fill.fore_color.rgb = LIGHT_GRAY
shape.line.color.rgb = ACCENT_ORANGE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA-cheng (+4 lines)", font_size=18, bold=True,
              color=ACCENT_ORANGE, alignment=PP_ALIGN.CENTER, space_after=Pt(8))

add_paragraph(tf, "DSD offset, same structure:", font_size=14,
              bold=True, color=DARK_GRAY, space_after=Pt(6))

add_paragraph(tf, "fn score_matvec_mult() {", font_size=12,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "+ var xk_off = kv_head_id*dim*seq;", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "+ right_dsd = @set_dsd_base_addr(", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "+     right_dsd, ptr_XKCache);", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "+ right_dsd = @increment_dsd_offset(", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "+     right_dsd, xk_off, f16);", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "  Kt=dim_p_pe; Nt=seq_len_p_pe;", font_size=12,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "- vecmat_computation();", font_size=12,
              font_name="Consolas", color=ACCENT_RED, space_after=Pt(2))
add_paragraph(tf, "+ vecmat_computation_dsd();", font_size=12,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "  all_reduce_bsz_seq_len(); // same", font_size=12,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "}", font_size=12, font_name="Consolas", color=DARK_GRAY, space_after=Pt(8))

add_paragraph(tf, "Minimal: offset into replicated cache,", font_size=13,
              bold=True, color=ACCENT_GREEN, space_after=Pt(2))
add_paragraph(tf, "swap to DSD-snapshot GEMV variant", font_size=13,
              bold=True, color=ACCENT_GREEN)

add_slide_number(slide, 7)


# ─────────────────────────────────────────────
# SLIDE 8: Communication Library Delta
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "COMMUNICATION")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Communication Library: The Biggest Delta",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

comm_rows = [
    ["Component", "Base", "WSE-3-GQA", "Delta", "GQA-cheng", "Delta"],
    ["comm_pe.csl", "1035 lines", "1708 lines", "+673 (+65%)", "1086 lines", "+51 (+5%)"],
    ["comm_layout.csl", "405 lines", "575 lines", "+170 (+42%)", "406 lines", "+1 (+0%)"],
    ["Reduce functions", "7", "~12", "+5", "~8", "+1"],
    ["Axis modes", "2 (X, Y)", "3 (X, Y, KV-head)", "+1", "2 (X, Y)", "0"],
]

table = add_table(slide, len(comm_rows), 6,
                  Inches(0.6), Inches(1.85), Inches(12.1), Inches(2.3))

col_widths = [Inches(2.2), Inches(1.7), Inches(1.9), Inches(2.0), Inches(1.9), Inches(2.4)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for row_idx, row_data in enumerate(comm_rows):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        if row_idx == 0:
            style_table_cell(cell, cell_text, font_size=13, bold=True,
                             color=WHITE, bg_color=TABLE_HEADER_BG,
                             alignment=PP_ALIGN.CENTER)
        else:
            bg = TABLE_ALT_BG if row_idx % 2 == 0 else WHITE
            is_bold = col_idx == 0
            clr = DARK_GRAY
            if col_idx == 3 and "+" in cell_text and cell_text != "+1":
                clr = ACCENT_RED
            elif col_idx == 5 and (cell_text == "+1 (+0%)" or cell_text == "0" or "+51" in cell_text):
                clr = ACCENT_GREEN
            style_table_cell(cell, cell_text, font_size=12, bold=is_bold,
                             color=clr, bg_color=bg,
                             alignment=PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT)

# WSE-3-GQA new functions
tf = add_rich_textbox(slide, Inches(0.6), Inches(4.4), Inches(5.8), Inches(2.8))
add_paragraph(tf, "WSE-3-GQA: New Communication Functions", font_size=16, bold=True,
              color=ACCENT_BLUE, space_after=Pt(6))
add_bullet(tf, "all_reduce_bsz_g()  -- group-extended softmax sum", font_size=13, color=DARK_GRAY)
add_bullet(tf, "all_reduceMax_bsz_g()  -- group-extended softmax max", font_size=13, color=DARK_GRAY)
add_bullet(tf, "all_reduce_bsz_g_var_seq_len_kv_head_scoped()", font_size=13, color=DARK_GRAY)
add_paragraph(tf, "    Head-scoped variable-length score reduce", font_size=12, color=MED_GRAY)
add_bullet(tf, "all_reduce_bsz_var_seq_len_strided()", font_size=13, color=DARK_GRAY)
add_bullet(tf, "reconfig_allreduce_axis(3)  -- new KV-head-scoped X-axis", font_size=13, color=DARK_GRAY)

tf = add_rich_textbox(slide, Inches(6.9), Inches(4.4), Inches(5.8), Inches(2.8))
add_paragraph(tf, "WSE-3-GQA-cheng: Minimal Additions", font_size=16, bold=True,
              color=ACCENT_ORANGE, space_after=Pt(6))
add_bullet(tf, "all_reduce_bsz_g()  -- placeholder (~50 lines)", font_size=13, color=DARK_GRAY)
add_paragraph(tf, "", font_size=10)
add_paragraph(tf, "Communication library essentially unchanged from base.", font_size=14,
              color=MED_GRAY, space_after=Pt(6))
add_paragraph(tf, "No new axis modes, no new routing topology.", font_size=14,
              color=MED_GRAY, space_after=Pt(6))
add_paragraph(tf, "comm_layout.csl delta: +1 line (+0.2%)", font_size=14,
              bold=True, color=ACCENT_GREEN)

add_slide_number(slide, 8)


# ─────────────────────────────────────────────
# SLIDE 9: Decode Pipeline Comparison
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "PIPELINE")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Decode Pipeline: Three-Way Comparison",
            font_size=30, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

pipeline_rows = [
    ["Step", "Base (WSE-3)", "WSE-3-GQA", "WSE-3-GQA-cheng"],
    ["1-2", "RMSNorm, Q proj", "Same", "Same"],
    ["3-4", "K proj [dim->dim]\nV proj [dim->dim]", "K proj [dim->KV_DIM]\nV proj [dim->KV_DIM]", "K proj [dim->dim, OFFSET]\nV proj [dim->dim, OFFSET]"],
    ["5", "Reduce QKV [3*dim]", "Reduce QKV [dim+2*kv]", "Reduce QKV [3*dim]"],
    ["6-8", "Reconfig(0), RoPE Q+K", "Reconfig(0), RoPE OFF", "Reconfig(0), RoPE Q+K"],
    ["9-10", "--", "PROCESS_KV()\nRECONFIG(3) head-scoped", "--"],
    ["11", "Score [dim x seq]", "Score MULTI-GROUP\nVAR-LEN, HEAD-SCOPED", "Score [dim x seq, OFFSET]"],
    ["12-13", "Reconfig(1), Softmax\n[bsz scalars]", "Reconfig(1), Softmax\n[groups*bsz scalars, VAR]", "Reconfig(1), Softmax\n[bsz scalars]"],
    ["14", "Output [seq x dim]", "Output MULTI-GROUP\nVAR-LEN", "Output [seq x dim, OFFSET]"],
    ["15-21", "O-proj, Residual, FFN\n(up+gate, reduce, SILU)", "Same pipeline\n(SILU bug inherited)", "Same pipeline\n(SILU bug FIXED)"],
]

table = add_table(slide, len(pipeline_rows), 4,
                  Inches(0.4), Inches(1.85), Inches(12.5), Inches(5.0))

col_widths = [Inches(1.0), Inches(3.3), Inches(4.5), Inches(3.7)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for row_idx, row_data in enumerate(pipeline_rows):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        if row_idx == 0:
            style_table_cell(cell, cell_text, font_size=13, bold=True,
                             color=WHITE, bg_color=TABLE_HEADER_BG,
                             alignment=PP_ALIGN.CENTER)
        else:
            bg = TABLE_ALT_BG if row_idx % 2 == 0 else WHITE
            clr = DARK_GRAY
            is_bold = col_idx == 0
            # Highlight changed rows
            if row_idx in [5, 6, 7, 8] and col_idx == 2:
                clr = ACCENT_BLUE
                is_bold = True
            elif row_idx in [2, 3, 6, 8] and col_idx == 3:
                clr = ACCENT_ORANGE
            elif "FIXED" in cell_text:
                clr = ACCENT_GREEN
                is_bold = True
            elif "OFF" in cell_text or "inherited" in cell_text:
                clr = ACCENT_RED
            style_table_cell(cell, cell_text, font_size=11, bold=is_bold,
                             color=clr, bg_color=bg,
                             alignment=PP_ALIGN.CENTER if col_idx == 0 else PP_ALIGN.LEFT)

add_slide_number(slide, 9)


# ─────────────────────────────────────────────
# SLIDE 10: Code Complexity Summary
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "COMPLEXITY")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Code Complexity Summary",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

complexity_rows = [
    ["Metric", "Base", "WSE-3-GQA", "Delta", "GQA-cheng", "Delta"],
    ["decode.csl", "965", "1282", "+317 (+33%)", "942", "-23 (-2%)"],
    ["comm_pe.csl", "1035", "1708", "+673 (+65%)", "1086", "+51 (+5%)"],
    ["comm_layout.csl", "405", "575", "+170 (+42%)", "406", "+1 (+0%)"],
    ["layout.csl", "114", "131", "+17 (+15%)", "120", "+6 (+5%)"],
    ["Total CSL", "2519", "3696", "+1177 (+47%)", "2554", "+35 (+1%)"],
    ["launch_sim.py", "336", "~2000", "+1664 (5x)", "~100", "-236"],
    ["GEMV variants", "1", "3", "+2", "2", "+1"],
    ["Reduce functions", "7", "~12", "+5", "~8", "+1"],
    ["Debug exports", "1", "7", "+6", "1", "0"],
]

table = add_table(slide, len(complexity_rows), 6,
                  Inches(0.6), Inches(1.85), Inches(12.1), Inches(4.5))

col_widths = [Inches(2.2), Inches(1.5), Inches(1.8), Inches(2.3), Inches(1.8), Inches(2.5)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for row_idx, row_data in enumerate(complexity_rows):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        if row_idx == 0:
            style_table_cell(cell, cell_text, font_size=13, bold=True,
                             color=WHITE, bg_color=TABLE_HEADER_BG,
                             alignment=PP_ALIGN.CENTER)
        elif row_idx == 5:  # Total row
            style_table_cell(cell, cell_text, font_size=13, bold=True,
                             color=BLACK, bg_color=RGBColor(0xE8, 0xEE, 0xF5),
                             alignment=PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT)
        else:
            bg = TABLE_ALT_BG if row_idx % 2 == 0 else WHITE
            is_bold = col_idx == 0
            clr = DARK_GRAY
            if col_idx == 3 and "+" in cell_text:
                clr = ACCENT_RED
            elif col_idx == 5:
                if "-" in cell_text and cell_text != "0":
                    clr = ACCENT_GREEN
                elif cell_text in ["0", "+1", "+1 (+0%)"]:
                    clr = ACCENT_GREEN
                elif "+51" in cell_text or "+6" in cell_text:
                    clr = ACCENT_GREEN
            style_table_cell(cell, cell_text, font_size=12, bold=is_bold,
                             color=clr, bg_color=bg,
                             alignment=PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT)

add_slide_number(slide, 10)


# ─────────────────────────────────────────────
# SLIDE 11: Key Findings
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "KEY FINDINGS")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Key Findings & Recommendations",
            font_size=32, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

# Left column - Findings
tf = add_rich_textbox(slide, Inches(0.6), Inches(1.9), Inches(6.0), Inches(5.2))

findings = [
    ("1. Radically different delta sizes",
     "WSE-3-GQA: +47% CSL code (heavy refactor)\nWSE-3-GQA-cheng: +1% CSL code (lightweight patch)"),
    ("2. Communication is the biggest delta",
     "comm_pe.csl: +673 lines (+65%) for WSE-3-GQA\nvs only +51 lines (+5%) for GQA-cheng"),
    ("3. Memory efficiency vs simplicity",
     "Dimensional: K/V shrinks by GQA ratio (optimal)\nReplicated: K/V grows by n_kv_heads (costly)"),
    ("4. WSE-3-GQA adds orthogonal features",
     "KV cache management, variable-length attention,\nprefill support, debug instrumentation"),
    ("5. Bug handling diverges",
     "GQA-cheng fixes SILU + DOWN bugs from base\nWSE-3-GQA inherits them unchanged"),
    ("6. RoPE regression",
     "WSE-3-GQA comments out RoPE (strided-DSD bug)\nGQA-cheng preserves it from the base"),
]

for title, desc in findings:
    add_paragraph(tf, title, font_size=15, bold=True, color=ACCENT_BLUE, space_after=Pt(2))
    add_paragraph(tf, desc, font_size=12, color=DARK_GRAY, space_after=Pt(10))

# Right column - Recommendations
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(7.0), Inches(1.9), Inches(5.8), Inches(5.2)
)
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xE8, 0xF5, 0xE9)
shape.line.color.rgb = ACCENT_GREEN

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "Recommendations", font_size=22, bold=True, color=ACCENT_GREEN,
              alignment=PP_ALIGN.CENTER, space_after=Pt(12))

add_paragraph(tf, "For Production:", font_size=16, bold=True, color=BLACK, space_after=Pt(4))
add_bullet(tf, "WSE-3-GQA's dimensional approach is the right long-term strategy", font_size=14, color=DARK_GRAY)
add_bullet(tf, "Memory efficiency scales with GQA ratio", font_size=14, color=DARK_GRAY)
add_bullet(tf, "Less communication volume is critical at scale", font_size=14, color=DARK_GRAY)

add_paragraph(tf, "", font_size=8)
add_paragraph(tf, "Backport from GQA-cheng:", font_size=16, bold=True, color=BLACK, space_after=Pt(4))
add_bullet(tf, "SILU iteration bug fix", font_size=14, color=DARK_GRAY)
add_bullet(tf, "DOWN weight shape fix", font_size=14, color=DARK_GRAY)
add_bullet(tf, "Re-enable RoPE (fix strided-DSD issue)", font_size=14, color=DARK_GRAY)

add_paragraph(tf, "", font_size=8)
add_paragraph(tf, "For Prototyping:", font_size=16, bold=True, color=BLACK, space_after=Pt(4))
add_bullet(tf, "GQA-cheng is valuable as a rapid prototyping path", font_size=14, color=DARK_GRAY)
add_bullet(tf, "Minimal delta makes it easy to test GQA correctness", font_size=14, color=DARK_GRAY)
add_bullet(tf, "Use before investing in full dimensional refactor", font_size=14, color=DARK_GRAY)

add_slide_number(slide, 11)


# ─────────────────────────────────────────────
# SLIDE 12: Host-Side Data Structure Handling
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "HOST-SIDE HANDLING")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Host-Side Data Structure: Q/O Permutation & KV Tiling",
            font_size=30, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

# Left: Q/O Permutation
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6), Inches(1.85), Inches(6.0), Inches(5.1)
)
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xEE, 0xF4, 0xFB)
shape.line.color.rgb = ACCENT_BLUE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "Q & O Weight Column Permutation", font_size=18, bold=True,
              color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER, space_after=Pt(6))
add_paragraph(tf, "Only needed in dimensional approach (WSE-3-GQA)", font_size=13,
              color=MED_GRAY, alignment=PP_ALIGN.CENTER, space_after=Pt(8))

add_paragraph(tf, "Problem:", font_size=15, bold=True, color=ACCENT_RED, space_after=Pt(3))
add_paragraph(tf, "Standard Llama Q layout: columns ordered by Q-head.", font_size=13,
              color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "But on-chip, each PE's Q output must be split into", font_size=13,
              color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "gqa_group_size contiguous kv_dim_p_pe slices,", font_size=13,
              bold=True, color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "each aligned with the KV-head's cache on that PE.", font_size=13,
              color=DARK_GRAY, space_after=Pt(8))

add_paragraph(tf, "Solution: offline column reorder", font_size=15, bold=True,
              color=ACCENT_GREEN, space_after=Pt(3))
add_paragraph(tf, "for _h in range(n_heads):", font_size=11,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(1))
add_paragraph(tf, "  _kv_head = _h // gqa_group_size", font_size=11,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(1))
add_paragraph(tf, "  _g = _h % gqa_group_size", font_size=11,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(1))
add_paragraph(tf, "  for _s in range(pes_p_kv_head):", font_size=11,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(1))
add_paragraph(tf, "    old = _h*head_dim + _s*kv_dim_p_pe", font_size=11,
              font_name="Consolas", color=ACCENT_RED, space_after=Pt(1))
add_paragraph(tf, "    new = _kv*pes_p_kv*dim + _s*dim + _g*kv", font_size=11,
              font_name="Consolas", color=ACCENT_GREEN, space_after=Pt(1))
add_paragraph(tf, "    W_Q_perm[:,new:new+kv] = W_Q[:,old:old+kv]", font_size=11,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(6))

add_paragraph(tf, "Result per PE: Q output = [grp0|grp1|...|grp_{G-1}]", font_size=13,
              bold=True, color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "Each grp is kv_dim_p_pe contiguous elements", font_size=13,
              color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "O weight rows reordered identically (transpose)", font_size=13,
              color=MED_GRAY)

# Right: KV Cache Tiling + Why
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(7.0), Inches(1.85), Inches(5.8), Inches(5.1)
)
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xFD, 0xF2, 0xE9)
shape.line.color.rgb = ACCENT_ORANGE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "KV Cache: Interleaved Tiling", font_size=18, bold=True,
              color=ACCENT_ORANGE, alignment=PP_ALIGN.CENTER, space_after=Pt(6))
add_paragraph(tf, "Custom tiling for asymmetric + batched + partial-fill", font_size=13,
              color=MED_GRAY, alignment=PP_ALIGN.CENTER, space_after=Pt(8))

add_paragraph(tf, "Base: simple reshape-transpose", font_size=14, bold=True,
              color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "XKCache[dim, seq_len] -> reshape(P,..,P,..) -> tile", font_size=12,
              font_name="Consolas", color=MED_GRAY, space_after=Pt(8))

add_paragraph(tf, "Dimensional: 3 new requirements", font_size=14, bold=True,
              color=ACCENT_ORANGE, space_after=Pt(3))
add_bullet(tf, "Asymmetric: kv_dim_p_pe per PE (not dim_p_pe)", font_size=13, color=DARK_GRAY)
add_bullet(tf, "Batched: [bsz, kv_dim, max_seq_len] with batch dim", font_size=13, color=DARK_GRAY)
add_bullet(tf, "Partial fill: only prefill_len tokens, rest zeroed", font_size=13, color=DARK_GRAY)
add_bullet(tf, "Interleaved: token t -> PE row py = t % P", font_size=13, color=DARK_GRAY)

add_paragraph(tf, "", font_size=6)
add_paragraph(tf, "On-chip layout per PE:", font_size=14, bold=True, color=DARK_GRAY, space_after=Pt(3))
add_paragraph(tf, "K: [bsz, kv_dim_p_pe, max_seq_len_p_pe]", font_size=12,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "   feature-major (for Q @ K^T)", font_size=12, color=MED_GRAY, space_after=Pt(2))
add_paragraph(tf, "V: [bsz, max_seq_len_p_pe, kv_dim_p_pe]", font_size=12,
              font_name="Consolas", color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "   sequence-major (for score @ V)", font_size=12, color=MED_GRAY, space_after=Pt(6))

add_paragraph(tf, "Why interleaving: PE row py owns tokens", font_size=13,
              bold=True, color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "{py, py+P, py+2P, ...} so Y-axis softmax", font_size=13,
              color=DARK_GRAY, space_after=Pt(2))
add_paragraph(tf, "reduction operates on distinct tokens", font_size=13, color=DARK_GRAY)

add_slide_number(slide, 12)


# ─────────────────────────────────────────────
# SLIDE 13: Host-Side Memcpy & Complexity
# ─────────────────────────────────────────────
slide = add_blank_slide()
add_section_badge(slide, "HOST-SIDE HANDLING")
add_textbox(slide, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
            "Host-Side: Transfer Sizes & Complexity Trade-off",
            font_size=30, bold=True, color=BLACK)
add_accent_line(slide, Inches(0.6), Inches(1.6), Inches(3), ACCENT_BLUE)

# Transfer size table
xfer_rows = [
    ["Symbol", "Base count_per_pe", "Dimensional count_per_pe", "Ratio (4:1 GQA)"],
    ["Q_weight", "dim_p_pe x dim_p_pe", "dim_p_pe x dim_p_pe", "1x (same)"],
    ["K_weight", "dim_p_pe x dim_p_pe", "dim_p_pe x kv_dim_p_pe", "1/4 (4x smaller)"],
    ["V_weight", "dim_p_pe x dim_p_pe", "dim_p_pe x kv_dim_p_pe", "1/4 (4x smaller)"],
    ["O_weight", "dim_p_pe x dim_p_pe", "dim_p_pe x dim_p_pe", "1x (same)"],
    ["XKCache", "dim_p_pe x seq_len_p_pe", "bsz x kv_dim_p_pe x max_seq", "varies"],
    ["XVCache", "seq_len_p_pe x dim_p_pe", "bsz x max_seq x kv_dim_p_pe", "varies"],
]

table = add_table(slide, len(xfer_rows), 4,
                  Inches(0.6), Inches(1.85), Inches(12.1), Inches(3.0))

col_widths = [Inches(1.8), Inches(3.5), Inches(3.8), Inches(3.0)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for row_idx, row_data in enumerate(xfer_rows):
    for col_idx, cell_text in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        if row_idx == 0:
            style_table_cell(cell, cell_text, font_size=13, bold=True,
                             color=WHITE, bg_color=TABLE_HEADER_BG,
                             alignment=PP_ALIGN.CENTER)
        else:
            bg = TABLE_ALT_BG if row_idx % 2 == 0 else WHITE
            clr = DARK_GRAY
            if col_idx == 3 and "smaller" in cell_text:
                clr = ACCENT_GREEN
            style_table_cell(cell, cell_text, font_size=12,
                             bold=(col_idx == 0 or (col_idx == 3 and "smaller" in cell_text)),
                             color=clr, bg_color=bg,
                             alignment=PP_ALIGN.CENTER if col_idx >= 2 else PP_ALIGN.LEFT)

# Complexity trade-off summary
shape = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.9)
)
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(0xFD, 0xF9, 0xED)
shape.line.color.rgb = ACCENT_ORANGE

tf = shape.text_frame
tf.word_wrap = True
add_paragraph(tf, "  Key Insight: Where the Complexity Lives", font_size=17, bold=True,
              color=ACCENT_ORANGE, space_after=Pt(6))

# Use a mini table-like layout with bullets
add_paragraph(tf, "  Host-side preprocessing (paid once at init):", font_size=14,
              bold=True, color=DARK_GRAY, space_after=Pt(3))
add_paragraph(tf, "    Generalized asymmetric tile_weight_row()  |  Offline Q/O column permutation (4 nested loops)  |  Custom interleaved KV cache tiling",
              font_size=13, color=MED_GRAY, space_after=Pt(8))

add_paragraph(tf, "  On-chip benefits (continuous during inference):", font_size=14,
              bold=True, color=ACCENT_GREEN, space_after=Pt(3))
add_paragraph(tf, "    4x smaller K/V buffers per PE  |  50% less QKV communication volume  |  No runtime head-mapping logic  |  Contiguous group slicing in score loops",
              font_size=13, color=DARK_GRAY)

add_slide_number(slide, 13)


# ─────────────────────────────────────────────
# SLIDE 14: Summary / Closing
# ─────────────────────────────────────────────
slide = add_blank_slide()

add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_BLUE)

add_textbox(slide, Inches(1), Inches(2.0), Inches(11.3), Inches(1.0),
            "Summary",
            font_size=40, bold=True, color=BLACK, alignment=PP_ALIGN.CENTER)

add_accent_line(slide, Inches(5.0), Inches(2.9), Inches(3.3), ACCENT_BLUE)

# Two-column summary boxes
shape_l = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(1.5), Inches(3.3), Inches(4.8), Inches(2.5)
)
shape_l.fill.solid()
shape_l.fill.fore_color.rgb = RGBColor(0xEE, 0xF4, 0xFB)
shape_l.line.color.rgb = ACCENT_BLUE

tf = shape_l.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA", font_size=24, bold=True, color=ACCENT_BLUE,
              alignment=PP_ALIGN.CENTER, space_after=Pt(8))
add_paragraph(tf, "Heavy refactor (+47% code)", font_size=16, color=DARK_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(4))
add_paragraph(tf, "Optimal memory efficiency", font_size=16, color=DARK_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(4))
add_paragraph(tf, "Right for production", font_size=18, bold=True, color=ACCENT_GREEN,
              alignment=PP_ALIGN.CENTER)

shape_r = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(7.0), Inches(3.3), Inches(4.8), Inches(2.5)
)
shape_r.fill.solid()
shape_r.fill.fore_color.rgb = RGBColor(0xFD, 0xF2, 0xE9)
shape_r.line.color.rgb = ACCENT_ORANGE

tf = shape_r.text_frame
tf.word_wrap = True
add_paragraph(tf, "WSE-3-GQA-cheng", font_size=24, bold=True, color=ACCENT_ORANGE,
              alignment=PP_ALIGN.CENTER, space_after=Pt(8))
add_paragraph(tf, "Lightweight patch (+1% code)", font_size=16, color=DARK_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(4))
add_paragraph(tf, "Preserves base structure", font_size=16, color=DARK_GRAY,
              alignment=PP_ALIGN.CENTER, space_after=Pt(4))
add_paragraph(tf, "Right for prototyping", font_size=18, bold=True, color=ACCENT_GREEN,
              alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(1), Inches(6.2), Inches(11.3), Inches(0.6),
            "Both approaches have value -- combine GQA-cheng's bug fixes with WSE-3-GQA's architecture",
            font_size=16, color=MED_GRAY, alignment=PP_ALIGN.CENTER)


# Save
output_path = "/mnt/raid0sata2/lexu/lexu/WaferLLM/Decode/GQA_Implementation_Comparison.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
